from abc import ABC, abstractmethod
import io
import re
from typing import TYPE_CHECKING, Any, Dict, List, Optional

if TYPE_CHECKING:
    from services.ingestion.app.pdf_layout import PdfLayoutArtifact


def _trim_bounds(text: str, start: int, end: int) -> tuple[int, int]:
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return start, end

class ParserException(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code

class ParsedSection:
    def __init__(
        self,
        title: str,
        content: str,
        page_number: Optional[int] = None,
        start_char: int = 0,
        end_char: Optional[int] = None,
    ):
        self.title = title
        self.content = content
        self.page_number = page_number
        self.start_char = start_char
        self.end_char = start_char + len(content) if end_char is None else end_char

class ParsedDocument:
    def __init__(
        self,
        raw_text: str,
        sections: List[ParsedSection],
        metadata: Optional[Dict[str, Any]] = None,
        pdf_layout: Optional["PdfLayoutArtifact"] = None,
    ):
        self.raw_text = raw_text
        self.sections = sections
        self.metadata = metadata or {}
        self.pdf_layout = pdf_layout
        for section in sections:
            if not 0 <= section.start_char <= section.end_char <= len(raw_text):
                raise ValueError("parsed section is outside the parsed artifact")
            if raw_text[section.start_char:section.end_char] != section.content:
                raise ValueError("parsed section content does not match its artifact span")

class BaseParser(ABC):
    @abstractmethod
    def parse(self, content_bytes: bytes, filename: str) -> ParsedDocument:
        pass

class TextParser(BaseParser):
    def parse(self, content_bytes: bytes, filename: str) -> ParsedDocument:
        if b'\x00' in content_bytes:
            raise ParserException("Text documents cannot contain binary NUL bytes.", status_code=415)
        try:
            text = content_bytes.decode('utf-8', errors='strict')
        except UnicodeDecodeError as error:
            raise ParserException("Text documents must be valid UTF-8.", status_code=415) from error
        if not text.strip():
            raise ParserException("The uploaded document contains no text.", status_code=422)
        sections = [
            ParsedSection(
                title="Document Content",
                content=text,
                page_number=1,
                start_char=0,
                end_char=len(text),
            )
        ]
        return ParsedDocument(raw_text=text, sections=sections, metadata={"filename": filename, "format": "text", "language": "en"})

class MarkdownParser(BaseParser):
    def parse(self, content_bytes: bytes, filename: str) -> ParsedDocument:
        if b'\x00' in content_bytes:
            raise ParserException("Markdown documents cannot contain binary NUL bytes.", status_code=415)
        try:
            text = content_bytes.decode('utf-8', errors='strict')
        except UnicodeDecodeError as error:
            raise ParserException("Markdown documents must be valid UTF-8.", status_code=415) from error
        if not text.strip():
            raise ParserException("The uploaded document contains no text.", status_code=422)
        sections: List[ParsedSection] = []
        
        # Split by Markdown headings (# Header)
        heading_pattern = re.compile(r'^(#{1,6})\s+(.+)$', re.MULTILINE)
        matches = list(heading_pattern.finditer(text))
        
        if not matches:
            sections.append(
                ParsedSection(
                    title="Main Content",
                    content=text,
                    page_number=1,
                    start_char=0,
                    end_char=len(text),
                )
            )
        else:
            preamble_start, preamble_end = _trim_bounds(text, 0, matches[0].start())
            if preamble_start < preamble_end:
                sections.append(
                    ParsedSection(
                        title="Preamble",
                        content=text[preamble_start:preamble_end],
                        page_number=1,
                        start_char=preamble_start,
                        end_char=preamble_end,
                    )
                )
            for i, match in enumerate(matches):
                title = match.group(2).strip()
                start_pos = match.end()
                end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(text)
                start_pos, end_pos = _trim_bounds(text, start_pos, end_pos)
                if start_pos < end_pos:
                    sections.append(
                        ParsedSection(
                            title=title,
                            content=text[start_pos:end_pos],
                            page_number=1,
                            start_char=start_pos,
                            end_char=end_pos,
                        )
                    )
        
        return ParsedDocument(raw_text=text, sections=sections, metadata={"filename": filename, "format": "markdown", "language": "en"})

class PDFParser(BaseParser):
    def parse(self, content_bytes: bytes, filename: str) -> ParsedDocument:
        if not content_bytes.startswith(b"%PDF-"):
            raise ParserException("The uploaded file is not a valid PDF document.", status_code=415)
        try:
            try:
                from app.pdf_layout import extract_native_pdf_layout
            except ModuleNotFoundError:
                from services.ingestion.app.pdf_layout import extract_native_pdf_layout

            layout = extract_native_pdf_layout(content_bytes, filename)
            sections: List[ParsedSection] = []
            pages_with_text = 0
            source_page_count = len(layout.pages)
            for page in layout.pages:
                page_text = layout.raw_text[page.parsed_start:page.parsed_end]
                if page_text.strip():
                    pages_with_text += 1
                sections.append(
                    ParsedSection(
                        title=f"Page {page.page_index + 1}",
                        content=page_text,
                        page_number=page.page_index + 1,
                        start_char=page.parsed_start,
                        end_char=page.parsed_end,
                    )
                )
            combined_text = layout.raw_text
            if not combined_text.strip():
                raise ParserException(
                    "The PDF contains no extractable text. Image-only PDFs require OCR, which is not enabled yet.",
                    status_code=422,
                )
            return ParsedDocument(
                raw_text=combined_text,
                sections=sections,
                metadata={
                    "filename": filename,
                    "format": "pdf",
                    "page_count": source_page_count,
                    "pages_with_text": pages_with_text,
                    "page_join_contract": "physical_pages_form_feed:v1",
                    "text_extraction_contract": "rawdict_sorted_spans_linefeed:v1",
                    "layout_schema_version": 1,
                    "language": "en",
                },
                pdf_layout=layout,
            )
        except ValueError as error:
            if str(error) == "encrypted_pdf":
                raise ParserException(
                    "PDF is password-protected or encrypted. "
                    "Please upload an unprotected document.",
                    status_code=400,
                ) from error
            if str(error) == "layout_limit":
                raise ParserException(
                    "The PDF is too structurally complex for bounded spatial extraction.",
                    status_code=413,
                ) from error
            raise ParserException(
                "The PDF could not be converted into a valid spatial text artifact.",
                status_code=422,
            ) from error
        except ParserException:
            raise
        except Exception as error:
            raise ParserException("The PDF could not be parsed as a valid document.", status_code=422) from error

class DocxParser(BaseParser):
    def parse(self, content_bytes: bytes, filename: str) -> ParsedDocument:
        if not content_bytes.startswith(b"PK"):
            raise ParserException("The uploaded file is not a valid DOCX document.", status_code=415)
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            paragraphs: List[tuple[str, bool]] = []
            for p in doc.paragraphs:
                text = p.text
                if not text.strip():
                    continue
                is_heading = bool(p.style and p.style.name.startswith("Heading"))
                paragraphs.append((text, is_heading))

            combined_parts: List[str] = []
            paragraph_spans: List[tuple[int, int, str, bool]] = []
            artifact_length = 0
            for text, is_heading in paragraphs:
                if combined_parts:
                    combined_parts.append("\n\n")
                    artifact_length += 2
                start = artifact_length
                combined_parts.append(text)
                artifact_length += len(text)
                paragraph_spans.append((start, artifact_length, text, is_heading))
            combined = "".join(combined_parts)
            if not combined.strip():
                raise ParserException("The DOCX document contains no extractable text.", status_code=422)

            sections: List[ParsedSection] = []
            current_title = "Introduction"
            body_start: Optional[int] = None
            body_end: Optional[int] = None

            def append_section() -> None:
                nonlocal body_start, body_end
                if body_start is not None and body_end is not None:
                    sections.append(
                        ParsedSection(
                            title=current_title,
                            content=combined[body_start:body_end],
                            page_number=1,
                            start_char=body_start,
                            end_char=body_end,
                        )
                    )
                body_start = None
                body_end = None

            for start, end, text, is_heading in paragraph_spans:
                if is_heading:
                    append_section()
                    current_title = text
                else:
                    if body_start is None:
                        body_start = start
                    body_end = end
            append_section()

            if not sections:
                sections.append(
                    ParsedSection(
                        title=current_title,
                        content=combined,
                        page_number=1,
                        start_char=0,
                        end_char=len(combined),
                    )
                )
            return ParsedDocument(raw_text=combined, sections=sections, metadata={"filename": filename, "format": "docx", "language": "en"})
        except ParserException:
            raise
        except Exception as error:
            raise ParserException("The DOCX file could not be parsed as a valid document.", status_code=422) from error

class ParserFactory:
    _parsers = {
        "application/pdf": PDFParser,
        "text/markdown": MarkdownParser,
        "text/plain": TextParser,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocxParser,
    }

    @classmethod
    def get_parser(cls, mime_type: str, filename: str = "") -> BaseParser:
        clean_mime = mime_type.split(';')[0].strip().lower()
        ext = filename.split('.')[-1].lower() if '.' in filename else ""
        if ext == "pdf":
            return PDFParser()
        elif ext in ["md", "markdown"]:
            return MarkdownParser()
        elif ext == "docx":
            return DocxParser()
        elif ext in ["txt", "text", "log", "json", "yaml", "yml"]:
            return TextParser()

        if ext:
            raise ParserException(
                f"Unsupported file extension '.{ext}'. Supported: PDF, Markdown, DOCX, TXT.",
                status_code=415,
            )

        if clean_mime in cls._parsers:
            return cls._parsers[clean_mime]()

        raise ParserException(f"Unsupported format '{ext or mime_type}'. Supported: PDF, Markdown, DOCX, TXT.", status_code=415)
